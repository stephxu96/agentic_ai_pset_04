import sys, json, glob, re, zipfile
from pptx import Presentation
SP, OUT = sys.argv[1], sys.argv[2]; d=json.load(open(f"{SP}/scenes.json"))
for f in sorted(glob.glob(OUT+"/*.pptx")):
    n=re.search(r"piece(\d\d)",f).group(1); p=Presentation(f); ok=0
    for i,s in enumerate(p.slides):
        notes=s.notes_slide.notes_text_frame.text
        if notes.startswith(d[n][i]["narration"]): ok+=1
        else: print("MISMATCH",n,i+1)
    z=zipfile.ZipFile(f); adv=[int(re.search(r'advTm="(\d+)"',z.read(f"ppt/slides/slide{k+1}.xml").decode()).group(1)) for k in range(len(p.slides))]
    print(n, len(p.slides),"slides; notes exact:",ok,"; auto-advance total %.1fs"%(sum(adv)/1000))
